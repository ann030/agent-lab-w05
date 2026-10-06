import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors
    C_NAVY = RGBColor(15, 23, 42)      # #0f172a
    C_BLUE = RGBColor(37, 99, 235)     # #2563eb
    C_LIGHT_BG = RGBColor(248, 250, 252) # #f8fafc
    C_WHITE = RGBColor(255, 255, 255)
    C_CARD_BG = RGBColor(255, 255, 255)
    C_CARD_BORDER = RGBColor(226, 232, 240)
    C_TEXT_MAIN = RGBColor(30, 41, 59) # #1e293b
    C_TEXT_MUTED = RGBColor(100, 116, 139) # #64748b
    C_ACCENT_GREEN = RGBColor(16, 185, 129)
    C_ACCENT_RED = RGBColor(239, 68, 68)

    def add_bg(slide, color=C_LIGHT_BG):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title, category="NDHU AGENT LAB · PRACTICE REPO"):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = C_BLUE

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = C_NAVY

    def add_card(slide, left, top, width, height, title, items, tag="", tag_color=C_BLUE):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD_BG
        card.line.color.rgb = C_CARD_BORDER
        card.line.width = Pt(1.5)

        tb = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.2), Inches(width - 0.5), Inches(height - 0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        if tag:
            p0.text = f"[{tag}]  {title}"
        else:
            p0.text = title
        p0.font.size = Pt(16)
        p0.font.bold = True
        p0.font.color.rgb = C_NAVY
        p0.space_after = Pt(12)

        for item in items:
            p = tf.add_paragraph()
            p.text = f"•  {item}"
            p.font.size = Pt(12.5)
            p.font.color.rgb = C_TEXT_MAIN
            p.space_after = Pt(6)

    # ---------------- SLIDE 1: COVER ----------------
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1, C_NAVY)

    tb = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11), Inches(3.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "國立東華大學 · 校園 AI Agent 實作成果報告"
    p0.font.size = Pt(16)
    p0.font.bold = True
    p0.font.color.rgb = RGBColor(96, 165, 250) # Light blue
    p0.space_after = Pt(16)

    p1 = tf.add_paragraph()
    p1.text = "NDHU Campus Agent Lab (`agent-lab-w05`)"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = C_WHITE
    p1.space_after = Pt(20)

    p2 = tf.add_paragraph()
    p2.text = "從日常小事出發：看懂計畫、驗收成果、堅守邊界、退回危險做法\n學生儲存庫：https://github.com/ann030/agent-lab-w05"
    p2.font.size = Pt(15)
    p2.font.color.rgb = RGBColor(203, 213, 225)
    p2.space_after = Pt(28)

    p3 = tf.add_paragraph()
    p3.text = "執行環境：Antigravity Agent + Git / GitHub | 完成項目：Task A, B, C, D 及實作紀錄"
    p3.font.size = Pt(13)
    p3.font.bold = True
    p3.font.color.rgb = C_ACCENT_GREEN

    # ---------------- SLIDE 2: OVERVIEW & OBJECTIVES ----------------
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)
    add_header(s2, "核心目標：培養 AI 協作的四大關鍵判斷力")

    cards_s2 = [
        ("1. 看懂計畫", ["在 AI 動作前仔細審查其白話計畫", "確認讀取與輸出路徑，不盲目授權", "辨認破壞性或越權操作"], "PLAN"),
        ("2. 檢查成果", ["使用客觀指標（如 Hash 雜湊比對）", "逐項執行驗收測試案例，不靠直覺", "抽查檔案副本與原檔是否完整一致"], "VERIFY"),
        ("3. 辨認權限", ["嚴守最小權限原則與工作區隔離", "原始輸入資料一律唯讀，禁止刪改", "未經授權絕不連網、上傳或公開"], "BOUNDARY"),
        ("4. 退回做法", ["發現高風險指令時果斷拒絕執行", "指出問題核心並給予合規替代方案", "保持人為主導、AI 輔助的清晰分工"], "REJECT")
    ]
    col_w = 2.7
    gap = 0.3
    start_x = 0.8
    for i, (title, items, tag) in enumerate(cards_s2):
        add_card(s2, start_x + i * (col_w + gap), 1.9, col_w, 4.8, title, items, tag)

    # ---------------- SLIDE 3: TASK A ----------------
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)
    add_header(s3, "Task A｜社團檔案安全整理與盤點 (01-club-files)")

    items_a1 = [
        "面對 12 個檔名雜亂、有多個 final 的社團文件",
        "原檔唯讀原則：input/ 目錄 100% 保持不動",
        "拒絕以檔名斷定：proposal_final 與 proposal_final2 內容完全不同（室外 vs 室內），皆標註未定案，不可自行刪除",
        "全部保留副本，分類至 proposals, meetings, logistics, promotion"
    ]
    items_a2 = [
        "MD5 雜湊抓出完全相同檔案：announcement 與 equipment 兩組",
        "相同內容亦安全保留副本，交由幹部人工確認",
        "產出 output/manifest.json（12 筆物件對照來源、目的與理由）",
        "產出 output/report.md 完整記錄盤點與待決策問題",
        "Git Commit: `A: organize club files`"
    ]
    add_card(s3, 0.8, 1.9, 5.7, 4.8, "整理原則與防護機制", items_a1, "安全準則", C_BLUE)
    add_card(s3, 6.8, 1.9, 5.7, 4.8, "雜湊驗收與結構化產出", items_a2, "驗收成果", C_ACCENT_GREEN)

    # ---------------- SLIDE 4: TASK B ----------------
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)
    add_header(s4, "Task B｜課間活動挑選器：單頁應用與 v1 到 v2 迭代")

    items_b1 = [
        "純前端單頁架構：內嵌 12 筆活動資料，雙擊即可離線開啟",
        "多條件嚴格篩選：地點（室內/室外/不限）、可用時間（15/30/60m）、強度（低/中/不限）",
        "無符合條件時清楚顯示提示，絕不擅自偷改條件",
        "歷史紀錄上限 5 筆（最新在前），提供篩選重設與中英雙語切換",
        "標明「教學模擬活動，不是校方公告」"
    ]
    items_b2 = [
        "現況 (v1)：僅能使用滑鼠點擊「幫我選」按鈕",
        "改進需求 (v2)：新增鍵盤快捷鍵（Space 空白鍵 / Enter 快速抽卡）",
        "體驗優化：按鍵微縮動態視覺回饋，多語系快捷鍵提示",
        "防呆設計：焦點在下拉選單時自動略過，不干擾原生選單操作",
        "Git Commits: `B v1: activity picker` ➔ `B v2: add shortcuts`"
    ]
    add_card(s4, 0.8, 1.9, 5.7, 4.8, "版本 1 (v1) 功能規格", items_b1, "v1 基礎功能", C_BLUE)
    add_card(s4, 6.8, 1.9, 5.7, 4.8, "版本 2 (v2) 迭代改進", items_b2, "v2 體驗升級", C_ACCENT_GREEN)

    # ---------------- SLIDE 5: TASK C ----------------
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5)
    add_header(s5, "Task C｜社團器材記錄清理：數據正規化與異常保留")

    items_c1 = [
        "文字欄位（item_id, name, status）去除前後贅餘空白",
        "數量欄位 (qty) 不套用去空白規則，嚴格保留原值",
        "狀態值標準化：可借/可出借/available 統一為 available；借出/borrowed 統一為 borrowed",
        "非標準狀態（如『待盤點』）標記為 unknown 並記錄於問題報告",
        "移除整列為空的無效列（列 6），10 筆原始資料正規化為 9 筆有效"
    ]
    items_c2 = [
        "缺漏數量（qty: ''）與負數（qty: -1）原值保留，不猜測補 0",
        "相同 item_id（EQ01, EQ02）全部保留，不擅自合併刪除",
        "指出衝突：EQ02 登記數量衝突（列 2 為 2，列 5 為 3）",
        "產出 normalized.json 與 issues.md 供實體盤點對照",
        "Git Commit: `C: clean equipment records`"
    ]
    add_card(s5, 0.8, 1.9, 5.7, 4.8, "清洗與正規化規則", items_c1, "資料清洗", C_BLUE)
    add_card(s5, 6.8, 1.9, 5.7, 4.8, "異常值保留與品質報告", items_c2, "客觀存證", C_ACCENT_GREEN)

    # ---------------- SLIDE 6: TASK D ----------------
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6)
    add_header(s6, "Task D｜安全防線：審查刻意寫錯的模擬計畫 (bad-plan.txt)")

    items_d1 = [
        "1. 越權存取 Downloads ➔ 違反最小權限，應限於指定題目資料夾",
        "2. 擅自刪除重複檔 ➔ 破壞不可逆，應保持唯讀並以報告列出",
        "3. 盲信 final2 為最新版 ➔ 檔名不可靠，不同版本應全數保留比對",
        "4. 缺資料隨意補合理值 ➔ 數據偽造，異常必須原樣保留並呈現",
        "5. 自動公開成果 ➔ 隱私與資料外洩風險，成果應僅存於本地 output"
    ]
    items_d2 = [
        "產出正式審查檔案：`practice/04-review/my-rejection.md`",
        "給 Agent 的明確指令範例：",
        "『這份計畫有多項不符規範，請停止並重新修正：",
        "  - 僅限操作指定目錄，嚴禁更動 Downloads",
        "  - 原始檔一律唯讀，禁止刪改，保留全部副本",
        "  - 不以 final 認定定稿，缺漏原樣保留",
        "  - 成果留於本地 output，嚴禁外連公開』",
        "Git Commit: `D: rejection`"
    ]
    add_card(s6, 0.8, 1.9, 5.7, 4.8, "圈出的 5 大高風險違規", items_d1, "問題診斷", C_ACCENT_RED)
    add_card(s6, 6.8, 1.9, 5.7, 4.8, "退回訊息與合規替代要求", items_d2, "退回範本", C_BLUE)

    # ---------------- SLIDE 7: SUMMARY ----------------
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7)
    add_header(s7, "實作歷程總覽與人機協作思維")

    items_s7_1 = [
        "• 815acee - record: learning record and screenshots",
        "• ba2a1f3 - C: clean equipment records",
        "• 97917bb - D: rejection",
        "• b452c63 - B v2: add Space and Enter shortcuts to pick activity",
        "• 7752faa - B v1: activity picker",
        "• 85941bd - A: organize club files",
        "每個階段皆獨立驗收、確認無誤後才 Commit & Push！"
    ]
    items_s7_2 = [
        "1. 人是核心決策者：邊界設定、安全審查與業務邏輯決定權在人",
        "2. AI 是高效率副駕駛：擅長大量結構化分類、程式碼生成與比對",
        "3. 凡事皆有憑據：透過 Git Commit 與學習紀錄留存客觀證據",
        "4. 安全與合規第一：拒絕黑盒，堅持看懂計畫才點同意"
    ]
    add_card(s7, 0.8, 1.9, 5.7, 4.8, "完整 Git 提交鏈 (Commit Chain)", items_s7_1, "GIT HISTORY", C_BLUE)
    add_card(s7, 6.8, 1.9, 5.7, 4.8, "結語：負責任的 AI 協作心態", items_s7_2, "KEY TAKEAWAY", C_ACCENT_GREEN)

    # Save presentation
    output_path = r"c:\Users\an108\OneDrive\桌面\大學作業\1006_AB\agent-lab-summary.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_deck()
